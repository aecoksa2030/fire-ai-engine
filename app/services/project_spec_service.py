import base64
import io
import os

import anthropic

from app.services.json_utils import clean_and_parse_json
from app.services.pricing import estimate_cost_usd

# Powers the standalone "Projects" page (app/routers/projects.py).
#
# Stage 1 of the estimation workflow: read a project's technical
# specification (PTS) and output its fire-protection SCOPE as a flat table
# — one row per (Area, Type of system required as per PTS) — exactly the
# layout of the estimation engineer's reference sheet ("Typical Expected
# output expected based on PTS.xlsx"). The engineer explicitly asked for
# this table ONLY, not a long narrative digest of the whole document: the
# output of this stage becomes the input of stage 2, together with the
# project's FEED package.
#
# Approved (engineer-reviewed) scope tables are saved to the reference
# library (ReferenceExample rows with report_type=REFERENCE_REPORT_TYPE,
# see app/routers/projects.py's /approve endpoint) and fed back into every
# new extraction as naming/granularity guidance — the same RAG approach
# the drawing analysis already uses (app/services/reference_library_service.py).

# report_type value used for this feature's rows in the shared
# reference_examples table. Keeps PTS scope examples fully separate from
# the drawing BOQ examples ("fire" / "security") without a schema change.
REFERENCE_REPORT_TYPE = "pts_scope"

# How many approved past scope tables to include per extraction, and how
# many of each one's rows. Enough to show the house style for Area/System
# naming and granularity without bloating every request.
REFERENCE_EXAMPLE_LIMIT = 3
REFERENCE_ROWS_PER_EXAMPLE = 80

SYSTEM_INSTRUCTION = (
    "You are a senior fire-protection estimation engineer. You are given a"
    " project technical specification (PTS) document for a substation or"
    " industrial facility. Your only task is to extract the fire-protection"
    " SCOPE of the project: every area the PTS names, and every type of"
    " system the PTS requires in that area. This scope table is the first"
    " step of the estimation workflow and is later used together with the"
    " project's FEED package, so it must be complete and must follow the"
    " document exactly — use only what the PTS states, never general"
    " knowledge or typical practice."
)

EXTRACTION_PROMPT = (
    "Read the entire attached PTS document — every page, every table and"
    " every coverage statement (including detector-type tables,"
    " extinguisher tables, clean agent / sprinkler / hydrant coverage"
    " clauses, life-safety equipment lists, passive protection clauses,"
    " and the SAS/ECC alarm lists, which show which buildings, rooms and"
    " equipment are protected).\n\n"
    "Return a single JSON object of exactly this shape:\n"
    '{"scope": [{"area": string, "system": string}]}\n\n'
    "Rules:\n"
    "1. One row per (area, system) pair. An area that needs three systems"
    " gets three rows.\n"
    "2. Keep all rows of the same area together, one after the other."
    " Order areas as they appear in the PTS.\n"
    "3. area = the area / room / building / equipment name as the PTS"
    ' writes it (e.g. "Control Room", "33kV Switchgear Hall", "Cable'
    ' Tunnels", "Power Transformers"). When the PTS lists several areas'
    " together in one sentence or table cell, split them into separate"
    " areas. Keep the PTS's voltage levels and qualifiers; do not merge"
    " areas the PTS treats separately, and do not invent areas the PTS"
    " does not mention. Keep conditional wording in the name when the PTS"
    ' uses it (e.g. "Cable Chutes (where applicable)").\n'
    "4. system = the type of system only — no sizes, ratings, agent types,"
    " detector types, standards or quantities. Use these names when they"
    ' fit: "Fire Detection and Alarm", "Clean Agent", "Sprinkler",'
    ' "Extinguishers", "Fire Hydrant", "Fire Hose Cabinet", "Water Spray",'
    ' "Foam", "SCBA", "First Aid", "Eye Wash and Safety Shower", "Passive'
    ' Fire Protection". If the PTS requires a type not in this list, name'
    " it in the same short way. Smoke/heat/beam/duct/gas detectors, manual"
    " call points and repeater or annunciator panels are all part of"
    ' "Fire Detection and Alarm" for that area, not separate systems.\n'
    "5. Include every system the PTS ties to an area — detection,"
    " suppression, portable extinguishers, life-safety equipment and"
    " passive fire protection. Do not add a system to an area unless the"
    " PTS requires it there.\n"
    "6. No duplicate rows.\n\n"
    "Return ONLY the raw JSON object — no markdown fences, no commentary."
)

# Large enough that a scope table (typically a few hundred short rows at
# most) can never realistically be cut off. Sent through the streaming API
# (see _stream_message) because the Anthropic SDK refuses a plain
# non-streaming call once max_tokens is high enough that it estimates the
# reply could take more than ~10 minutes.
EXTRACT_MAX_TOKENS = 32000


def build_reference_block(examples: list) -> str:
  """Turns approved past scope tables (ReferenceExample rows) into a short
  text block prepended to the extraction prompt. Returns "" when the
  library has no approved PTS scopes yet."""
  if not examples:
    return ""

  lines = [
      "REFERENCE — scope tables from OTHER projects that an estimation"
      " engineer reviewed and approved. Use them ONLY as a guide to the"
      " expected naming style and level of detail for Area and System."
      " They are different projects: do NOT copy any area or system from"
      " them unless THIS PTS requires it.",
  ]
  for i, example in enumerate(examples, start=1):
    rows = normalize_scope_rows((example.approved_result or {}).get("scope"))
    lines.append(f"\nApproved example {i} ({example.tag}):")
    for row in rows[:REFERENCE_ROWS_PER_EXAMPLE]:
      lines.append(f"  {row['area']} | {row['system']}")
    if len(rows) > REFERENCE_ROWS_PER_EXAMPLE:
      lines.append(f"  ... ({len(rows) - REFERENCE_ROWS_PER_EXAMPLE} more rows)")

  return "\n".join(lines) + "\n\n---\n\n"


def normalize_scope_rows(rows) -> list:
  """Cleans a list of {"area", "system"} rows: strips whitespace, drops
  empty/malformed rows and duplicates (case-insensitive), and keeps every
  area's rows together in order of the area's first appearance — so the
  output always matches the reference sheet's grouping even if the model
  (or a manual edit) scattered one area's rows."""
  if not isinstance(rows, list):
    return []

  grouped = {}  # area key -> {"area": display name, "systems": [...]}
  order = []
  seen = set()
  for row in rows:
    if not isinstance(row, dict):
      continue
    area = str(row.get("area") or "").strip()
    system = str(row.get("system") or "").strip()
    if not area or not system:
      continue
    area_key = area.lower()
    pair_key = (area_key, system.lower())
    if pair_key in seen:
      continue
    seen.add(pair_key)
    if area_key not in grouped:
      grouped[area_key] = {"area": area, "systems": []}
      order.append(area_key)
    grouped[area_key]["systems"].append(system)

  return [
      {"area": grouped[k]["area"], "system": s}
      for k in order
      for s in grouped[k]["systems"]
  ]


def normalize_scope_data(data) -> dict:
  """Guarantees the stored/returned shape {"scope": [...], "truncated":
  bool, ...}. Records saved before this format existed (the old long
  digest) come back with an empty scope instead of crashing the page."""
  if not isinstance(data, dict):
    data = {}
  out = {
      "scope": normalize_scope_rows(data.get("scope")),
      "truncated": bool(data.get("truncated", False)),
  }
  for key in ("api_usage", "reference_examples_used"):
    if key in data:
      out[key] = data[key]
  return out


def build_scope_xlsx(rows: list) -> bytes:
  """Writes the scope table in exactly the reference sheet's layout: one
  sheet, header row at C4:E4 ("SL #", "Area", "Type of system required as
  per PTS") with the same light-blue fill (Office theme accent 1, 80%
  lighter), thin borders on every cell, Calibri 11, data from row 5. The
  Area/System columns are only widened when a value would otherwise be
  cut off."""
  from openpyxl import Workbook
  from openpyxl.styles import Border, PatternFill, Side
  from openpyxl.styles.colors import Color

  wb = Workbook()
  ws = wb.active
  ws.title = "Sheet1"

  thin = Side(style="thin", color=Color(auto=True))
  border = Border(left=thin, right=thin, top=thin, bottom=thin)
  header_fill = PatternFill(
      patternType="solid",
      fgColor=Color(theme=4, tint=0.7999816888943144),
      bgColor=Color(indexed=64),
  )

  headers = ["SL #", "Area", "Type of system required as per PTS"]
  for col, text in zip(("C", "D", "E"), headers):
    cell = ws[f"{col}4"]
    cell.value = text
    cell.fill = header_fill
    cell.border = border

  for i, row in enumerate(rows, start=1):
    r = 4 + i
    for col, value in zip(("C", "D", "E"), (i, row["area"], row["system"])):
      cell = ws[f"{col}{r}"]
      cell.value = value
      cell.border = border

  area_width = max([16.42578125] + [len(r["area"]) * 1.1 + 2 for r in rows])
  system_width = max([32.42578125] + [len(r["system"]) * 1.1 + 2 for r in rows])
  ws.column_dimensions["C"].width = 4.28515625
  ws.column_dimensions["D"].width = area_width
  ws.column_dimensions["E"].width = system_width

  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue()


class ProjectSpecService:

  def __init__(self):
    self.client = anthropic.Anthropic(
        api_key=os.environ.get("ANTHROPIC_API_KEY"),
        max_retries=5,
        # Well past the SDK's 10-minute default, so a long document can
        # never be cut short by a client-side clock.
        timeout=1800.0,
    )
    # Same env var / default as claude_service.py and claude_chat_service.py.
    self.model = os.environ.get("CLAUDE_MODEL", "claude-fable-5")

  def _stream_message(self, *, system: str, messages: list, max_tokens: int):
    """One messages call through the SDK's streaming helper; returns the
    same accumulated Message object a plain .create() call would."""
    with self.client.messages.stream(
        model=self.model, max_tokens=max_tokens, system=system, messages=messages,
    ) as stream:
      for _ in stream.text_stream:
        pass
      return stream.get_final_message()

  def extract(self, pdf_bytes: bytes, reference_block: str = "") -> dict:
    """Blocking call — the router runs it in a ThreadPoolExecutor so a long
    document doesn't freeze the single-process server for other users.
    Returns {"scope": [...], "truncated": bool, "api_usage": {...}}, plus
    "error" (a message string) if the call failed."""
    file_base64 = base64.b64encode(pdf_bytes).decode("utf-8")

    try:
      response = self._stream_message(
          system=SYSTEM_INSTRUCTION,
          messages=[{
              "role": "user",
              "content": [
                  {
                      "type": "document",
                      "source": {
                          "type": "base64",
                          "media_type": "application/pdf",
                          "data": file_base64,
                      },
                  },
                  {"type": "text", "text": reference_block + EXTRACTION_PROMPT},
              ],
          }],
          max_tokens=EXTRACT_MAX_TOKENS,
      )

      result_text = next(
          (b.text for b in response.content if hasattr(b, "text")), ""
      ).strip()
      parsed = clean_and_parse_json(result_text) if result_text else {}

      data = normalize_scope_data({"scope": parsed.get("scope")})
      # A flag, not a message — the frontend shows the warning in the
      # user's selected UI language.
      data["truncated"] = response.stop_reason == "max_tokens"
      data["api_usage"] = {
          "provider": "claude",
          "model": response.model,
          "input_tokens": response.usage.input_tokens,
          "output_tokens": response.usage.output_tokens,
          "total_cost_usd": estimate_cost_usd(
              response.model,
              response.usage.input_tokens,
              response.usage.output_tokens,
          ),
      }
      return data

    except Exception as e:
      return {
          "scope": [],
          "truncated": False,
          "error": str(e),
          "api_usage": {"provider": "claude", "model": self.model, "input_tokens": 0,
                        "output_tokens": 0, "total_cost_usd": 0.0},
      }
